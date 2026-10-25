"""
===============================================================================
 MY EXPENSE TRACKER  (Streamlit single-file app)
===============================================================================

A personal, single-user expense tracker. Log daily spends, see where the
money goes, and track each pay cycle against your budget. Currency: PKR.

-------------------------------------------------------------------------------
 HOW THIS FILE IS ORGANIZED  (search for the ===== banners to jump around)
-------------------------------------------------------------------------------
   1. CONFIG            - categories, colours, constants
   2. THEME             - the look (CSS + colours). Pure styling, no logic.
   3. DATABASE          - connect once, create tables, read/write helpers
   4. DATA ACCESS       - simple functions: add/edit/delete/load expenses
   5. MONEY & DATES     - formatting PKR, and the pay-cycle / salary-date maths
   6. EXCEL EXPORT      - build a nicely formatted .xlsx file
   7. LOAD STATE        - read settings + categories + the chosen period
   8. UI: HEADER        - the always-visible "how much can I spend" banner
   9. UI: TABS          - Add, Insights, Notes, Manage, Categories, Export
-------------------------------------------------------------------------------
 Storage: uses a permanent Postgres database if a connection string is given
 (st.secrets["DATABASE_URL"] on Streamlit Cloud, or the DATABASE_URL env var).
 With no connection string it falls back to a local SQLite file for testing.

 Run locally:   pip install -r requirements.txt   then   streamlit run app.py
===============================================================================
"""

import os
from io import BytesIO
from datetime import date, datetime, timedelta
from calendar import monthrange

import pandas as pd
import altair as alt
import streamlit as st


# =============================================================================
# 1. CONFIG
# =============================================================================

# The categories that always exist. The user can add more in the Categories tab.
BUILTIN_CATEGORIES = [
    "InDrive to office", "InDrive back home",
    "Breakfast", "Lunch", "Dinner",
    "Rent", "Internet bill", "Mobile package",
    "Miscellaneous",
]

# Every category belongs to a "group". Groups power the Fixed-vs-flexible donut.
BUILTIN_GROUPS = {
    "InDrive to office": "Commute", "InDrive back home": "Commute",
    "Breakfast": "Food", "Lunch": "Food", "Dinner": "Food",
    "Rent": "Fixed monthly", "Internet bill": "Fixed monthly", "Mobile package": "Fixed monthly",
    "Miscellaneous": "Other",
}

# Groups a user can choose from when creating a custom category.
GROUP_OPTIONS = ["Commute", "Food", "Fixed monthly", "Entertainment", "Shopping", "Health", "Other"]

# The quick-add buttons on the Add tab: (button label, real category name).
QUICK_PICKS = [
    ("🚗 To office", "InDrive to office"),
    ("🏠 Back home", "InDrive back home"),
    ("🍳 Breakfast", "Breakfast"),
    ("🍽️ Lunch", "Lunch"),
    ("🌙 Dinner", "Dinner"),
]

CURRENCY = "PKR"
DEFAULT_SALARY_DAY = 25  # salary normally arrives on the 25th

st.set_page_config(page_title="My Expense Tracker", page_icon="💸", layout="wide")


# =============================================================================
# 2. THEME  (pure styling - changing anything here cannot affect the numbers)
# =============================================================================

st.markdown(
    """
    <style>
      @import url('https://fonts.googleapis.com/css2?family=Poppins:wght@500;600;700&family=Inter:wght@400;500;600&display=swap');

      html, body, [class*="css"] { font-family: 'Inter', sans-serif; }
      .block-container { padding-top: 1.6rem; }

      /* Gradient app title */
      h1 {
          font-family: 'Poppins', sans-serif !important;
          font-weight: 700 !important;
          background: linear-gradient(90deg, #FF4FD8 0%, #6C5CE7 50%, #00D2FF 100%);
          -webkit-background-clip: text; -webkit-text-fill-color: transparent;
          background-clip: text; letter-spacing: -0.5px;
      }
      h2, h3, h4 { font-family: 'Poppins', sans-serif !important; letter-spacing: -0.3px; }

      /* Metric cards: rounded, soft shadow, lift on hover */
      div[data-testid="stMetricValue"] { font-size: 1.55rem; font-weight: 700; }
      div[data-testid="stMetric"] {
          border-radius: 16px; padding: 16px 18px;
          border: 1px solid rgba(255,255,255,0.08);
          box-shadow: 0 4px 18px rgba(0,0,0,0.25);
          transition: transform 0.15s ease;
      }
      div[data-testid="stMetric"]:hover { transform: translateY(-3px); }
      div[data-testid="column"]:nth-child(1) div[data-testid="stMetric"] {
          background: linear-gradient(135deg, rgba(255,79,216,0.18), rgba(255,79,216,0.04)); }
      div[data-testid="column"]:nth-child(2) div[data-testid="stMetric"] {
          background: linear-gradient(135deg, rgba(0,210,255,0.18), rgba(0,210,255,0.04)); }
      div[data-testid="column"]:nth-child(3) div[data-testid="stMetric"] {
          background: linear-gradient(135deg, rgba(0,184,148,0.18), rgba(0,184,148,0.04)); }
      div[data-testid="column"]:nth-child(4) div[data-testid="stMetric"] {
          background: linear-gradient(135deg, rgba(253,203,110,0.20), rgba(253,203,110,0.04)); }

      /* Primary buttons: gradient */
      button[kind="primary"] {
          background: linear-gradient(90deg, #6C5CE7, #FF4FD8) !important;
          border: none !important; border-radius: 10px !important; font-weight: 600 !important;
      }
      button[kind="primary"]:hover { filter: brightness(1.08); }

      /* Tabs + progress bar */
      button[data-baseweb="tab"] { font-weight: 600; }
      div[data-baseweb="tab-list"] { gap: 6px; }
      div[data-testid="stProgress"] > div > div > div > div {
          background: linear-gradient(90deg, #00D2FF, #6C5CE7, #FF4FD8) !important;
      }
    </style>
    """,
    unsafe_allow_html=True,
)


# =============================================================================
# 3. DATABASE  (one connection, reused -> the app stays fast)
# =============================================================================

def get_database_url():
    """Find the Postgres connection string, or return None to use local SQLite."""
    try:
        if "DATABASE_URL" in st.secrets:
            return st.secrets["DATABASE_URL"]
    except Exception:
        pass
    return os.environ.get("DATABASE_URL")


DATABASE_URL = get_database_url()
USE_POSTGRES = bool(DATABASE_URL)
SQLITE_PATH = "expenses.db"


@st.cache_resource
def _postgres_connection():
    """Open ONE Postgres connection and reuse it across reruns (this is the speed fix)."""
    import psycopg2
    conn = psycopg2.connect(
        DATABASE_URL, sslmode="require",
        keepalives=1, keepalives_idle=30, keepalives_interval=10, keepalives_count=5,
    )
    conn.autocommit = True
    return conn


def _placeholders(sql):
    """SQLite wants '?' placeholders; Postgres wants '%s'. We write '%s' everywhere."""
    return sql if USE_POSTGRES else sql.replace("%s", "?")


def run(sql, params=(), fetch=False):
    """
    Run one SQL statement.
      - fetch=False -> just execute (INSERT/UPDATE/DELETE/CREATE)
      - fetch=True  -> return the rows (SELECT)
    Postgres reuses the cached connection; SQLite opens a quick local connection.
    """
    sql = _placeholders(sql)
    if USE_POSTGRES:
        import psycopg2
        for attempt in range(2):  # retry once if the connection went stale
            try:
                cur = _postgres_connection().cursor()
                cur.execute(sql, params)
                rows = cur.fetchall() if fetch else None
                cur.close()
                return rows
            except (psycopg2.InterfaceError, psycopg2.OperationalError):
                _postgres_connection.clear()
                if attempt == 1:
                    raise
    else:
        import sqlite3
        conn = sqlite3.connect(SQLITE_PATH)
        try:
            cur = conn.cursor()
            cur.execute(sql, params)
            rows = cur.fetchall() if fetch else None
            conn.commit()
            return rows
        finally:
            conn.close()


def init_db():
    """Create the three tables the first time the app runs."""
    if USE_POSTGRES:
        run("""CREATE TABLE IF NOT EXISTS expenses (
                id SERIAL PRIMARY KEY, category TEXT NOT NULL, amount NUMERIC NOT NULL,
                description TEXT DEFAULT '', spent_on DATE NOT NULL,
                created_at TIMESTAMP DEFAULT NOW())""")
    else:
        run("""CREATE TABLE IF NOT EXISTS expenses (
                id INTEGER PRIMARY KEY AUTOINCREMENT, category TEXT NOT NULL, amount REAL NOT NULL,
                description TEXT DEFAULT '', spent_on TEXT NOT NULL,
                created_at TEXT DEFAULT CURRENT_TIMESTAMP)""")
    run("CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY, value TEXT)")
    run("CREATE TABLE IF NOT EXISTS categories (name TEXT PRIMARY KEY, grp TEXT)")


# Only create tables once per session, not on every little rerun.
if "db_ready" not in st.session_state:
    init_db()
    st.session_state["db_ready"] = True


# =============================================================================
# 4. DATA ACCESS  (small, named functions so the UI stays readable)
# =============================================================================

# ----- settings (salary, budget, view mode, per-month salary dates) -----
def load_settings():
    """Return all settings as a simple {key: value} dictionary."""
    rows = run("SELECT key, value FROM settings", fetch=True) or []
    return {k: v for k, v in rows}


def save_setting(key, value):
    if USE_POSTGRES:
        run("INSERT INTO settings (key, value) VALUES (%s, %s) "
            "ON CONFLICT (key) DO UPDATE SET value = EXCLUDED.value", (key, str(value)))
    else:
        run("INSERT OR REPLACE INTO settings (key, value) VALUES (%s, %s)", (key, str(value)))


# ----- custom categories -----
def load_custom_categories():
    """Return the user's own categories as a list of (name, group)."""
    rows = run("SELECT name, grp FROM categories ORDER BY name", fetch=True) or []
    return [(r[0], r[1]) for r in rows]


def save_category(name, group):
    if USE_POSTGRES:
        run("INSERT INTO categories (name, grp) VALUES (%s, %s) "
            "ON CONFLICT (name) DO UPDATE SET grp = EXCLUDED.grp", (name, group))
    else:
        run("INSERT OR REPLACE INTO categories (name, grp) VALUES (%s, %s)", (name, group))


def delete_category(name):
    run("DELETE FROM categories WHERE name = %s", (name,))


# ----- expenses: create / update / delete / read -----
def add_expense(category, amount, description, spent_on):
    run("INSERT INTO expenses (category, amount, description, spent_on) VALUES (%s, %s, %s, %s)",
        (category, float(amount), description or "", str(spent_on)))


def edit_expense(expense_id, category, amount, description, spent_on):
    run("UPDATE expenses SET category=%s, amount=%s, description=%s, spent_on=%s WHERE id=%s",
        (category, float(amount), description or "", str(spent_on), int(expense_id)))


def remove_expense(expense_id):
    run("DELETE FROM expenses WHERE id = %s", (int(expense_id),))


def load_expenses(start, end):
    """Return all expenses between two dates (inclusive) as a tidy DataFrame."""
    rows = run("SELECT id, category, amount, description, spent_on FROM expenses "
               "WHERE spent_on >= %s AND spent_on <= %s ORDER BY spent_on DESC, id DESC",
               (str(start), str(end)), fetch=True) or []
    df = pd.DataFrame(rows, columns=["id", "category", "amount", "description", "spent_on"])
    if not df.empty:
        df["amount"] = df["amount"].astype(float)
        df["spent_on"] = pd.to_datetime(df["spent_on"]).dt.date
        df["description"] = df["description"].fillna("")
    return df


def last_amount_for(category):
    """The most recent amount used for a category, so Add can pre-fill it. 0 if never used."""
    rows = run("SELECT amount FROM expenses WHERE category = %s ORDER BY spent_on DESC, id DESC LIMIT 1",
               (category,), fetch=True)
    return float(rows[0][0]) if rows else 0.0


# =============================================================================
# 5. MONEY & DATES
# =============================================================================

def money(x):
    """Format a number as 'PKR 12,345' (with thousands separators, no decimals)."""
    try:
        return f"{CURRENCY} {float(x):,.0f}"
    except Exception:
        return f"{CURRENCY} 0"


def prev_month(y, m):
    return (y - 1, 12) if m == 1 else (y, m - 1)


def next_month(y, m):
    return (y + 1, 1) if m == 12 else (y, m + 1)


def estimated_payday(year, month, normal_day):
    """
    The salary day for a month. If the normal day lands on a weekend we assume
    the salary came on the Friday before (e.g. Sun 25th -> Fri 23rd).
    """
    last = monthrange(year, month)[1]
    d = date(year, month, min(normal_day, last))
    weekday = d.weekday()  # Mon=0 ... Sat=5, Sun=6
    if weekday == 5:       # Saturday -> Friday
        d -= timedelta(days=1)
    elif weekday == 6:     # Sunday -> Friday
        d -= timedelta(days=2)
    return d


def payday(settings, year, month, normal_day):
    """The salary date for a month: a manual override if saved, else the estimate."""
    override = settings.get(f"payday_{year:04d}-{month:02d}")
    if override:
        try:
            return date.fromisoformat(override)
        except Exception:
            pass
    return estimated_payday(year, month, normal_day)


def next_payday(settings, normal_day, from_date):
    """The next salary date on or after a given date."""
    y, m = from_date.year, from_date.month
    for _ in range(4):
        d = payday(settings, y, m, normal_day)
        if d >= from_date:
            return d
        y, m = next_month(y, m)
    return payday(settings, y, m, normal_day)


def cycle_bounds(settings, normal_day, label_year, label_month):
    """
    A pay cycle labelled (year, month) runs from the PREVIOUS month's salary
    date up to the day before THIS month's salary date.
    Example: the 'November' cycle = 23 Oct (Oct salary) .. 24 Nov.
    """
    py, pm = prev_month(label_year, label_month)
    start = payday(settings, py, pm, normal_day)
    end = payday(settings, label_year, label_month, normal_day) - timedelta(days=1)
    return start, end


def current_cycle_label(settings, normal_day, today):
    """Which cycle are we in today? Once this month's salary arrives, we're in next month's cycle."""
    if today >= payday(settings, today.year, today.month, normal_day):
        return next_month(today.year, today.month)
    return (today.year, today.month)


# =============================================================================
# 6. EXCEL EXPORT
# =============================================================================

def build_excel(df, label, salary, budget, total_spent):
    """Return a formatted .xlsx (bytes) with an Expenses sheet and a Summary sheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="6C5CE7")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    title_font = Font(bold=True, size=14, color="1A1A2E")
    label_font = Font(bold=True, size=11)
    money_fmt = "#,##0"
    thin = Side(style="thin", color="D9D9E3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")

    # --- Sheet 1: Expenses ---
    ws = wb.active
    ws.title = "Expenses"
    ws["A1"] = f"Expenses - {label}"
    ws["A1"].font = title_font
    ws.merge_cells("A1:D1")
    for col, h in enumerate(["Date", "Category", "Amount (PKR)", "Description"], start=1):
        c = ws.cell(row=3, column=col, value=h)
        c.fill, c.font, c.alignment, c.border = header_fill, header_font, center, border

    r = 4
    for _, row in (df.sort_values("spent_on") if not df.empty else df).iterrows():
        ws.cell(r, 1, pd.to_datetime(row["spent_on"]).strftime("%Y-%m-%d")).border = border
        ws.cell(r, 2, row["category"]).border = border
        amt = ws.cell(r, 3, float(row["amount"]))
        amt.number_format, amt.border = money_fmt, border
        ws.cell(r, 4, row["description"]).border = border
        r += 1
    ws.cell(r, 2, "TOTAL").font = label_font
    total_cell = ws.cell(r, 3, float(total_spent))
    total_cell.number_format, total_cell.font = money_fmt, label_font
    for i, w in enumerate([14, 22, 16, 40], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"
    if r > 4:
        ws.auto_filter.ref = f"A3:D{r-1}"

    # --- Sheet 2: Summary ---
    ws2 = wb.create_sheet("Summary")
    ws2["A1"] = f"Summary - {label}"
    ws2["A1"].font = title_font
    ws2.merge_cells("A1:C1")
    rr = 3
    for lbl, val in [
        ("Salary", salary), ("Budget", budget), ("Total spent", total_spent),
        ("Left from budget", (budget - total_spent) if budget else 0),
        ("Left from salary", (salary - total_spent) if salary else 0),
    ]:
        ws2.cell(rr, 1, lbl).font = label_font
        vc = ws2.cell(rr, 2, float(val))
        vc.number_format = money_fmt
        rr += 1
    rr += 1
    for col, h in enumerate(["Category", "Amount (PKR)", "Share"], start=1):
        c = ws2.cell(rr, col, h)
        c.fill, c.font = header_fill, header_font
    rr += 1
    if not df.empty:
        for cat, amt in df.groupby("category")["amount"].sum().sort_values(ascending=False).items():
            ws2.cell(rr, 1, cat)
            ac = ws2.cell(rr, 2, float(amt))
            ac.number_format = money_fmt
            ws2.cell(rr, 3, f"{(amt / total_spent * 100) if total_spent else 0:.1f}%")
            rr += 1
    for i, w in enumerate([22, 16, 10], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


# =============================================================================
# 7. LOAD STATE  (read everything we need for this screen)
# =============================================================================

today = date.today()
settings = load_settings()
custom_cats = load_custom_categories()

# Build the live category list and group lookup (built-in + custom).
ALL_CATEGORIES = BUILTIN_CATEGORIES + [n for n, _ in custom_cats if n not in BUILTIN_CATEGORIES]
GROUPS = dict(BUILTIN_GROUPS)
for name, group in custom_cats:
    GROUPS[name] = group or "Other"

salary = float(settings.get("monthly_salary", 0) or 0)
budget = float(settings.get("monthly_budget", 0) or 0)
normal_day = int(settings.get("salary_day", DEFAULT_SALARY_DAY) or DEFAULT_SALARY_DAY)
use_cycle = settings.get("view_mode", "cycle") == "cycle"

st.title("💸 My Expense Tracker")

# ---- Period picker: which month/cycle are we looking at? ----
default_y, default_m = current_cycle_label(settings, normal_day, today) if use_cycle else (today.year, today.month)
years = list(range(today.year - 2, today.year + 2))
pick1, pick2, pick3 = st.columns([1, 1, 2])
with pick1:
    sel_year = st.selectbox("Year", years, index=years.index(default_y))
with pick2:
    sel_month = st.selectbox("Salary month" if use_cycle else "Month", list(range(1, 13)),
                             index=default_m - 1,
                             format_func=lambda m: datetime(2000, m, 1).strftime("%B"))
with pick3:
    choice = st.radio("View", ["Pay cycle", "Calendar month"], index=0 if use_cycle else 1,
                      horizontal=True,
                      help="Pay cycle = salary date to next salary date. Calendar = 1st to month end.")
    if (choice == "Pay cycle") != use_cycle:
        save_setting("view_mode", "cycle" if choice == "Pay cycle" else "calendar")
        st.rerun()

label = datetime(sel_year, sel_month, 1).strftime("%B %Y")

# Work out the date range for the chosen period.
if use_cycle:
    period_start, period_end = cycle_bounds(settings, normal_day, sel_year, sel_month)
    period_key = f"cycle_{sel_year:04d}-{sel_month:02d}"
else:
    last = monthrange(sel_year, sel_month)[1]
    period_start, period_end = date(sel_year, sel_month, 1), date(sel_year, sel_month, last)
    period_key = f"{sel_year:04d}-{sel_month:02d}"
range_text = f"{period_start.strftime('%d %b')} to {period_end.strftime('%d %b %Y')}"
viewing_now = period_start <= today <= period_end  # are we looking at the live period?

# Load this period's expenses and the key totals.
df = load_expenses(period_start, period_end)
total_spent = float(df["amount"].sum()) if not df.empty else 0.0
budget_left = budget - total_spent
salary_left = salary - total_spent


# ----------------------------------------------------------------------------
# SIDEBAR: the "set once" stuff (salary, budget, salary date)
# ----------------------------------------------------------------------------
with st.sidebar:
    st.header("⚙️ Setup")
    st.caption("Set these once. Saved automatically.")
    salary_in = st.number_input("Salary per cycle (PKR)", min_value=0, value=int(salary),
                                step=1000, format="%d")
    budget_in = st.number_input("Budget per cycle (PKR)", min_value=0, value=int(budget),
                                step=1000, format="%d", help="How much you plan to spend per cycle.")
    salday_in = st.number_input("Normal salary day", min_value=1, max_value=31, value=normal_day,
                                step=1, help="Usually 25. On a weekend the app uses the Friday before.")
    if st.button("💾 Save setup", use_container_width=True):
        save_setting("monthly_salary", salary_in)
        save_setting("monthly_budget", budget_in)
        save_setting("salary_day", salday_in)
        st.toast("Setup saved", icon="✅")
        st.rerun()

    st.divider()
    st.subheader("📅 Salary date")
    # In cycle view, the salary that STARTS this cycle came in the previous month.
    if use_cycle:
        st.caption(f"When the salary for the {label} cycle arrived. Change it if a holiday moved it.")
        py, pm = prev_month(sel_year, sel_month)
        default_pd = payday(settings, py, pm, normal_day)
        payday_key = f"payday_{py:04d}-{pm:02d}"
    else:
        st.caption(f"When the salary arrived in {label}. Change it if a holiday moved it.")
        default_pd = payday(settings, sel_year, sel_month, normal_day)
        payday_key = f"payday_{sel_year:04d}-{sel_month:02d}"

    actual_pd = st.date_input("Actual salary date", value=default_pd, key=f"pd_{period_key}")
    d1, d2 = st.columns(2)
    if d1.button("Save date", use_container_width=True):
        save_setting(payday_key, actual_pd.isoformat())
        st.toast("Salary date saved", icon="📅")
        st.rerun()
    if d2.button("Reset", use_container_width=True, help="Back to the automatic guess."):
        save_setting(payday_key, "")
        st.toast("Reset to automatic", icon="↩️")
        st.rerun()

    st.divider()
    st.caption("Storage: " + ("Postgres (permanent)" if USE_POSTGRES else "Local SQLite (testing only)"))


# =============================================================================
# 8. UI: HEADER  (always visible - answers "how much can I spend?")
# =============================================================================

st.markdown(f"#### {label} {'cycle' if use_cycle else 'month'}  ·  {range_text}")

# Gentle first-run nudge if salary/budget aren't set yet.
if salary == 0 and budget == 0:
    st.info("👋 Welcome! Start in the sidebar: set your **salary** and **budget** per cycle. "
            "Then log expenses in the ➕ Add tab.")

h1, h2, h3, h4 = st.columns(4)
h1.metric("Spent", money(total_spent))

if budget > 0:
    h2.metric("Budget left", money(budget_left), delta=f"of {money(budget)}", delta_color="off")
else:
    h2.metric("Budget left", "—", help="Set a budget in the sidebar.")

# Safe-to-spend per day: only meaningful while we're inside the live cycle.
if budget > 0 and viewing_now:
    days_left = max((period_end - today).days + 1, 1)
    safe_per_day = max(budget_left, 0) / days_left
    h3.metric("Safe to spend / day", money(safe_per_day), delta=f"{days_left} days left", delta_color="off")
else:
    h3.metric("Salary left", money(salary_left) if salary else "—")

nxt = next_payday(settings, normal_day, today)
days_to = (nxt - today).days
h4.metric("Next salary", nxt.strftime("%d %b"),
          delta=(f"in {days_to} days" if days_to > 0 else "today"), delta_color="off")

# One clear status line + a coloured progress bar.
if budget > 0:
    used = total_spent / budget
    st.progress(min(used, 1.0),
                text=f"Budget used: {money(total_spent)} / {money(budget)} ({used*100:.0f}%)")
    if total_spent > budget:
        st.error(f"You are over budget by {money(total_spent - budget)}.")
    elif used >= 0.85:
        st.warning(f"Careful: only {money(budget_left)} left this {'cycle' if use_cycle else 'month'}.")
    else:
        st.success(f"You have {money(budget_left)} left to spend.")

st.divider()


# =============================================================================
# 9. UI: TABS
# =============================================================================
tab_add, tab_insights, tab_notes, tab_manage, tab_cats, tab_export = st.tabs(
    ["➕ Add", "📊 Insights", "📝 Notes", "🧾 Manage", "🏷️ Categories", "⬇️ Export"]
)

# -----------------------------------------------------------------------------
# TAB: ADD  (the daily driver - fastest path to logging a spend)
# -----------------------------------------------------------------------------
with tab_add:
    st.subheader("Add an expense")

    # Quick-pick buttons set the category; the amount then pre-fills to last time.
    st.caption("Quick pick your usual, or choose any category below.")
    qcols = st.columns(len(QUICK_PICKS))
    for i, (btn_label, cat_name) in enumerate(QUICK_PICKS):
        if qcols[i].button(btn_label, use_container_width=True, key=f"quick_{cat_name}"):
            st.session_state["add_category"] = cat_name

    # Make sure the selectbox has a valid starting value.
    if st.session_state.get("add_category") not in ALL_CATEGORIES:
        st.session_state["add_category"] = ALL_CATEGORIES[0]

    a1, a2, a3 = st.columns([2, 1, 1])
    with a1:
        category = st.selectbox("Category", ALL_CATEGORIES, key="add_category")
    with a2:
        # Pre-fill with the last amount used for this category (own key per category).
        amount = st.number_input("Amount (PKR)", min_value=0, step=50, format="%d",
                                 value=int(last_amount_for(category)), key=f"amt_{category}")
    with a3:
        spent_on = st.date_input("Date", value=today, key="add_date")
    description = st.text_input("Description (optional)", key="add_desc",
                                placeholder="e.g. Movie with friends, medicine, gift")

    if st.button("Add expense", type="primary", use_container_width=True):
        if amount <= 0:
            st.warning("Enter an amount greater than 0.")
        else:
            add_expense(category, amount, description, spent_on)
            st.toast(f"Added {money(amount)} · {category}", icon="💸")
            st.rerun()

    # Immediate feedback: show the few most recent entries.
    if not df.empty:
        st.markdown("###### Recent")
        recent = df.head(5)[["spent_on", "category", "amount", "description"]].copy()
        recent["amount"] = recent["amount"].apply(money)
        recent.columns = ["Date", "Category", "Amount", "Description"]
        st.dataframe(recent, use_container_width=True, hide_index=True)

    st.caption("Need a category that isn't listed? Add it in the 🏷️ Categories tab.")

# -----------------------------------------------------------------------------
# TAB: INSIGHTS  (the charts)
# -----------------------------------------------------------------------------
with tab_insights:
    if df.empty:
        st.info("No expenses yet for this period. Add one in the ➕ Add tab and your charts appear here.")
    else:
        by_cat = df.groupby("category")["amount"].sum().sort_values(ascending=False).reset_index()
        by_cat.columns = ["Category", "Amount"]

        grouped = df.copy()
        grouped["Group"] = grouped["category"].map(lambda c: GROUPS.get(c, "Other"))
        by_group = grouped.groupby("Group")["amount"].sum().reset_index()
        by_group.columns = ["Group", "Amount"]

        c1, c2 = st.columns(2)
        with c1:
            st.markdown("**Spending by category**")
            st.altair_chart(
                alt.Chart(by_cat).mark_bar(cornerRadiusEnd=4, color="#6C5CE7").encode(
                    x=alt.X("Amount:Q", title="PKR", axis=alt.Axis(format=",.0f")),
                    y=alt.Y("Category:N", sort="-x", title=None),
                    tooltip=[alt.Tooltip("Category:N"), alt.Tooltip("Amount:Q", format=",.0f", title="PKR")],
                ).properties(height=300),
                use_container_width=True)
        with c2:
            st.markdown("**Daily spending**")
            by_day = df.groupby("spent_on")["amount"].sum().sort_index().reset_index()
            by_day.columns = ["Date", "Amount"]
            st.altair_chart(
                alt.Chart(by_day).mark_line(point=True, color="#00B894").encode(
                    x=alt.X("Date:T", title=None),
                    y=alt.Y("Amount:Q", title="PKR", axis=alt.Axis(format=",.0f")),
                    tooltip=[alt.Tooltip("Date:T"), alt.Tooltip("Amount:Q", format=",.0f", title="PKR")],
                ).properties(height=300),
                use_container_width=True)

        c3, c4 = st.columns(2)
        with c3:
            st.markdown("**Category share**")
            st.altair_chart(
                alt.Chart(by_cat).mark_arc().encode(
                    theta=alt.Theta("Amount:Q", stack=True),
                    color=alt.Color("Category:N", scale=alt.Scale(scheme="tableau20"),
                                    legend=alt.Legend(title=None, orient="bottom", columns=2)),
                    tooltip=[alt.Tooltip("Category:N"), alt.Tooltip("Amount:Q", format=",.0f", title="PKR")],
                ).properties(height=320),
                use_container_width=True)
        with c4:
            st.markdown("**Fixed vs flexible**")
            st.altair_chart(
                alt.Chart(by_group).mark_arc(innerRadius=70).encode(
                    theta=alt.Theta("Amount:Q", stack=True),
                    color=alt.Color("Group:N", scale=alt.Scale(scheme="set2"),
                                    legend=alt.Legend(title=None, orient="bottom")),
                    tooltip=[alt.Tooltip("Group:N"), alt.Tooltip("Amount:Q", format=",.0f", title="PKR")],
                ).properties(height=320),
                use_container_width=True)

        st.markdown("**Category breakdown**")
        table = by_cat.copy()
        table["Share"] = (table["Amount"] / total_spent * 100).round(1).astype(str) + "%"
        table["Amount"] = table["Amount"].apply(money)
        st.dataframe(table, use_container_width=True, hide_index=True)

# -----------------------------------------------------------------------------
# TAB: NOTES  (read every description, filterable)
# -----------------------------------------------------------------------------
with tab_notes:
    st.subheader("Notes & descriptions")
    if df.empty:
        st.info("No expenses this period yet.")
    else:
        with_notes = df[df["description"].str.strip() != ""]
        options = ["All categories"] + sorted(df["category"].unique().tolist())
        pick = st.selectbox("Filter by category", options)
        view = with_notes if pick == "All categories" else with_notes[with_notes["category"] == pick]

        if view.empty:
            st.info("No descriptions for this selection. Add a note when logging and it shows up here.")
        else:
            show = view.sort_values("spent_on", ascending=False)[
                ["spent_on", "category", "amount", "description"]].copy()
            show["amount"] = show["amount"].apply(money)
            show.columns = ["Date", "Category", "Amount", "Description"]
            st.dataframe(show, use_container_width=True, hide_index=True,
                         column_config={"Description": st.column_config.TextColumn(width="large")})
            st.caption(f"{len(view)} expense(s) with a description.")

# -----------------------------------------------------------------------------
# TAB: MANAGE  (edit or delete any row)
# -----------------------------------------------------------------------------
with tab_manage:
    st.subheader("Manage expenses")
    if df.empty:
        st.info("Nothing to manage for this period yet.")
    else:
        st.caption("Edit any cell. Use the row's trash icon to delete. Then press Save changes.")
        grid = df[["id", "spent_on", "category", "amount", "description"]].copy()
        edited = st.data_editor(
            grid, use_container_width=True, hide_index=True, num_rows="dynamic", key="editor",
            column_config={
                "id": None,
                "spent_on": st.column_config.DateColumn("Date", format="YYYY-MM-DD"),
                "category": st.column_config.SelectboxColumn("Category", options=ALL_CATEGORIES, required=True),
                "amount": st.column_config.NumberColumn("Amount (PKR)", min_value=0, step=50, format="%d"),
                "description": st.column_config.TextColumn("Description", width="large"),
            })

        if st.button("💾 Save changes", type="primary"):
            original = {int(r["id"]): r for _, r in df.iterrows()}
            seen, n_edit, n_add, n_del = set(), 0, 0, 0
            for _, row in edited.iterrows():
                rid = row.get("id")
                amount, category = row.get("amount"), row.get("category")
                spent, note = row.get("spent_on"), (row.get("description") or "")

                # a brand-new row (no id) that isn't filled in -> skip
                if pd.isna(rid) and (pd.isna(amount) or not category or not spent):
                    continue

                if pd.notna(rid):                      # existing row: update if changed
                    rid = int(rid)
                    seen.add(rid)
                    o = original.get(rid)
                    if o is not None and (str(o["spent_on"]) != str(spent) or o["category"] != category
                                          or float(o["amount"]) != float(amount or 0)
                                          or (o["description"] or "") != note):
                        edit_expense(rid, category, amount or 0, note, spent)
                        n_edit += 1
                elif amount and float(amount) > 0:     # new row: insert
                    add_expense(category, amount, note, spent)
                    n_add += 1

            for rid in original:                       # any row removed from the grid: delete
                if rid not in seen:
                    remove_expense(rid)
                    n_del += 1

            st.toast(f"Saved: {n_edit} edited, {n_add} added, {n_del} deleted", icon="✅")
            st.rerun()

# -----------------------------------------------------------------------------
# TAB: CATEGORIES  (add / remove your own)
# -----------------------------------------------------------------------------
with tab_cats:
    st.subheader("Your categories")
    st.caption("Built-in categories always exist. Add your own below, each in a group "
               "so it lands in the right place on the Fixed-vs-flexible chart.")

    with st.form("add_category_form", clear_on_submit=True):
        f1, f2 = st.columns([2, 1])
        new_name = f1.text_input("Category name", placeholder="e.g. Movie nights, InDrive (other), Gym")
        new_group = f2.selectbox("Group", GROUP_OPTIONS, index=GROUP_OPTIONS.index("Entertainment"))
        if st.form_submit_button("Add category", type="primary"):
            name = (new_name or "").strip()
            if not name:
                st.warning("Please type a category name.")
            elif name in ALL_CATEGORIES:
                st.warning("That category already exists.")
            else:
                save_category(name, new_group)
                st.toast(f"Added category: {name}", icon="🏷️")
                st.rerun()

    st.divider()
    st.markdown("**Built-in categories**")
    st.dataframe(pd.DataFrame([(c, BUILTIN_GROUPS[c]) for c in BUILTIN_CATEGORIES],
                              columns=["Category", "Group"]),
                 use_container_width=True, hide_index=True)

    st.markdown("**Your custom categories**")
    if not custom_cats:
        st.info("None yet. Add one above.")
    else:
        for name, group in custom_cats:
            col_n, col_g, col_d = st.columns([3, 2, 1])
            col_n.write(name)
            col_g.write(group or "Other")
            if col_d.button("Delete", key=f"delcat_{name}"):
                delete_category(name)
                st.toast(f"Deleted: {name}", icon="🗑️")
                st.rerun()
        st.caption("Deleting a category keeps any past expenses logged under it.")

# -----------------------------------------------------------------------------
# TAB: EXPORT  (download Excel or CSV)
# -----------------------------------------------------------------------------
with tab_export:
    st.subheader("Download your data")
    if df.empty:
        st.info("No data for this period to export yet.")
    else:
        st.download_button(
            "⬇️ Download Excel (formatted)",
            data=build_excel(df, label, salary, budget, total_spent),
            file_name=f"expenses_{period_key}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary", use_container_width=True)
        st.download_button(
            "⬇️ Download CSV (plain)",
            data=df[["spent_on", "category", "amount", "description"]].to_csv(index=False).encode("utf-8"),
            file_name=f"expenses_{period_key}.csv",
            mime="text/csv", use_container_width=True)
        st.caption("Excel has two sheets: a formatted Expenses list with a total, "
                   "and a Summary with salary, budget, and category breakdown.")
